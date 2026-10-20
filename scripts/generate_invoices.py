"""Generate the synthetic sample workbooks and the small test fixtures.

Outputs (see docs/data-design.md §9):
  data/sample/sample_clean.xlsx       clean 12-month sample (seeded)
  data/sample/sample_defects.xlsx     clean sample + Dataset A defect records
  data/sample/generation_report.json  actual sample profile vs. design targets
  tests/fixtures/cases/TC-*.xlsx      Dataset B isolated cases (TC-01 is .csv)
  tests/fixtures/worked_example.xlsx  hand-worked reference example (§6)

Expected results for the fixtures are hand-written in tests/fixtures/expected/
and are never produced by this script.

Amounts are integer halalas everywhere in this module; they are converted to
SAR only when a cell is written (§2).

Usage:
  .venv\\Scripts\\python.exe scripts\\generate_invoices.py [--seed N]
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "data" / "sample"
FIXTURES_DIR = ROOT / "tests" / "fixtures"

SEED = 20260930
RP_START = date(2025, 10, 1)
VD = date(2026, 9, 30)  # valuation date, also the last day of the reporting period
EXPORT_DATE = date(2026, 10, 15)  # no record is dated after this
VAT_PERCENT = 15
ISSUE_WEEKDAYS = {6, 0, 1, 2, 3}  # Sunday..Thursday (Monday == 0)
CREDIT_NOTE_SHARE = 0.04
CREDIT_REASONS = ("return", "damaged", "price_adjustment")

COLUMNS = {
    "customers": ["customer_id", "customer_name", "city", "segment", "payment_terms_days"],
    "invoices": ["invoice_id", "customer_id", "issue_date", "due_date", "currency",
                 "net_amount", "vat_amount", "total_amount"],
    "payments": ["payment_id", "invoice_id", "payment_date", "amount"],
    "credit_notes": ["credit_note_id", "invoice_id", "credit_date",
                     "net_amount", "vat_amount", "total_amount", "reason"],
}
AMOUNT_COLUMNS = {"net_amount", "vat_amount", "total_amount", "amount"}
DATE_COLUMNS = {"issue_date", "due_date", "payment_date", "credit_date"}

# tier -> (base invoices per month, min net SAR, max net SAR)
TIERS = {
    "key": (4.0, 10_000, 45_000),
    "medium": (2.0, 6_000, 25_000),
    "small": (1.5, 2_000, 9_000),
}
# Declared generation assumption, not a statistical finding (§3).
SEASONAL_FACTORS = {(2026, 1): 1.15, (2026, 2): 1.30, (2026, 3): 1.10,
                    (2026, 7): 0.85, (2026, 8): 0.85}
SEGMENT_LABELS = {"minimarket": "Mini Market", "grocery": "Grocery", "horeca": "Cafe"}
NAME_WORDS = ["Waha", "Nakheel", "Rawabi", "Sahab", "Bayan", "Qamar", "Reem", "Sanabel",
              "Ghadeer", "Fajr", "Najm", "Dana", "Yasmeen", "Shams", "Badr", "Sidra", "Wadi",
              "Thuraya", "Murooj", "Hilal", "Lamar", "Jood", "Nada", "Areen", "Ruba"]
# (tier, segment, payment terms, payment profile) for CUST-001 .. CUST-025 (§3 customer plan)
CUSTOMER_PLAN = [
    ("key", "minimarket", 60, "prompt"),
    ("key", "minimarket", 60, "late"),
    ("key", "minimarket", 90, "partial"),
    ("medium", "minimarket", 30, "prompt"),
    ("medium", "minimarket", 30, "late"),
    ("medium", "minimarket", 30, "prompt"),
    ("medium", "minimarket", 60, "late"),
    ("medium", "horeca", 60, "prompt"),
    ("medium", "horeca", 60, "late"),
    ("medium", "horeca", 60, "partial"),
    *[("small", "grocery", 30, "prompt")] * 6,
    *[("small", "grocery", 30, "late")] * 4,
    *[("small", "horeca", 30, "partial")] * 3,
    *[("small", "horeca", 30, "slow")] * 2,
]

TARGETS = {
    "gross_sales_sar": (6_500_000, 8_500_000),
    "invoice_count": (550, 650),
    "top1_net_sales_share": (0.15, 0.22),
    "top5_net_sales_share": (0.50, 0.66),
    "credit_note_ratio": (0.005, 0.01),
    "open_balance_months_of_sales": (1.5, 2.5),
    "over_90_share_of_open_balance": (0.03, 0.10),
    "empty_months": (0, 0),
}


@dataclass(frozen=True)
class Customer:
    customer_id: str
    name: str
    segment: str
    terms: int
    tier: str
    profile: str


@dataclass(frozen=True)
class Invoice:
    invoice_id: str
    customer_id: str
    issue_date: date
    due_date: date
    net: int
    vat: int
    total: int


@dataclass(frozen=True)
class Payment:
    payment_id: str
    invoice_id: str
    payment_date: date
    amount: int


@dataclass(frozen=True)
class CreditNote:
    credit_note_id: str
    invoice_id: str
    credit_date: date
    net: int
    vat: int
    total: int
    reason: str


@dataclass(frozen=True)
class Dataset:
    customers: tuple[Customer, ...]
    invoices: tuple[Invoice, ...]
    payments: tuple[Payment, ...]
    credit_notes: tuple[CreditNote, ...]


# ---------------------------------------------------------------- helpers

def vat_for(net: int) -> int:
    """VAT in halalas, ROUND_HALF_UP to the halala (net is positive)."""
    return (net * VAT_PERCENT + 50) // 100


def halalas_to_cell(halalas: int) -> float:
    """Nearest binary double to the SAR amount; the round trip is tested (§2)."""
    return halalas / 100


def rp_months() -> list[tuple[int, int]]:
    months, y, m = [], RP_START.year, RP_START.month
    for _ in range(12):
        months.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return months


def issue_days(year: int, month: int) -> list[date]:
    day, days = date(year, month, 1), []
    while day.month == month:
        if day.weekday() in ISSUE_WEEKDAYS:
            days.append(day)
        day += timedelta(days=1)
    return days


def poisson(rng: random.Random, lam: float) -> int:
    limit, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= limit:
            return k
        k += 1


def log_uniform_halalas(rng: random.Random, low_sar: int, high_sar: int) -> int:
    value = math.exp(rng.uniform(math.log(low_sar * 100), math.log(high_sar * 100)))
    return int(value + 0.5)


# ---------------------------------------------------------------- clean sample

def build_customers() -> tuple[Customer, ...]:
    customers = []
    for n, (tier, segment, terms, profile) in enumerate(CUSTOMER_PLAN, 1):
        name = f"{NAME_WORDS[n - 1]} {SEGMENT_LABELS[segment]} {n:02d}"
        customers.append(Customer(f"CUST-{n:03d}", name, segment, terms, tier, profile))
    return tuple(customers)


def payment_plan(rng: random.Random, profile: str, inv: Invoice, payable: int) -> list[tuple[date, int]]:
    """Payment dates and amounts for one invoice; amounts always sum to <= payable."""
    def on(offset: int) -> date:
        return max(inv.issue_date, inv.due_date + timedelta(days=offset))

    if profile == "prompt":
        return [(on(rng.randint(-10, 3)), payable)]
    if profile == "late":
        return [(on(rng.randint(10, 45)), payable)]
    if profile == "slow":
        return [(on(rng.randint(30, 120)), payable)] if rng.random() < 0.5 else []
    # partial: 2-3 instalments in whole riyals; the last one takes the remaining halalas
    first = int(payable * rng.uniform(0.4, 0.6)) // 100 * 100
    plan = [(on(rng.randint(-5, 10)), first)]
    rest = payable - first
    offsets = sorted(rng.sample(range(15, 61), rng.choice([1, 2])))
    for offset in offsets[:-1]:
        part = rest // 2 // 100 * 100
        plan.append((on(offset), part))
        rest -= part
    plan.append((on(offsets[-1]), rest))
    return plan


def generate_clean(seed: int = SEED) -> Dataset:
    rng = random.Random(seed)
    customers = build_customers()
    by_id = {c.customer_id: c for c in customers}

    drafts = []
    for c in customers:
        rate, low, high = TIERS[c.tier]
        for y, m in rp_months():
            days = issue_days(y, m)
            for _ in range(poisson(rng, rate * SEASONAL_FACTORS.get((y, m), 1.0))):
                drafts.append((rng.choice(days), c.customer_id, len(drafts),
                               log_uniform_halalas(rng, low, high)))
    drafts.sort()

    invoices = []
    for n, (issue, customer_id, _, net) in enumerate(drafts, 1):
        vat = vat_for(net)
        due = issue + timedelta(days=by_id[customer_id].terms)
        invoices.append(Invoice(f"INV-{n:05d}", customer_id, issue, due, net, vat, net + vat))

    credit_drafts, payment_drafts = [], []
    for inv in invoices:
        payable = inv.total
        if rng.random() < CREDIT_NOTE_SHARE:
            credit_date = inv.issue_date + timedelta(days=rng.randint(3, 20))
            cn_net = int(inv.net * rng.uniform(0.05, 0.30) + 0.5)
            reason = rng.choice(CREDIT_REASONS)
            if credit_date <= EXPORT_DATE:  # later credit notes are not in the export
                cn_vat = vat_for(cn_net)
                credit_drafts.append((credit_date, inv.invoice_id, cn_net, cn_vat, reason))
                payable -= cn_net + cn_vat
        for pay_date, amount in payment_plan(rng, by_id[inv.customer_id].profile, inv, payable):
            if pay_date <= EXPORT_DATE:  # later payments are not in the export
                payment_drafts.append((pay_date, inv.invoice_id, amount))

    credit_notes = tuple(
        CreditNote(f"CN-{n:05d}", invoice_id, d, net, vat, net + vat, reason)
        for n, (d, invoice_id, net, vat, reason) in enumerate(sorted(credit_drafts), 1)
    )
    payments = tuple(
        Payment(f"PAY-{n:05d}", invoice_id, d, amount)
        for n, (d, invoice_id, amount) in enumerate(sorted(payment_drafts), 1)
    )
    dataset = Dataset(customers, tuple(invoices), payments, credit_notes)
    problems = check_clean(dataset)
    if problems:
        raise ValueError("clean sample failed consistency checks:\n" + "\n".join(problems))
    return dataset


def check_clean(ds: Dataset) -> list[str]:
    """Consistency rules for the clean sample (§4 step 6). Returns problems found."""
    problems = []
    for name, ids in [("customers", [c.customer_id for c in ds.customers]),
                      ("invoices", [i.invoice_id for i in ds.invoices]),
                      ("payments", [p.payment_id for p in ds.payments]),
                      ("credit_notes", [c.credit_note_id for c in ds.credit_notes])]:
        if len(ids) != len(set(ids)):
            problems.append(f"{name}: duplicate IDs")

    terms = {c.customer_id: c.terms for c in ds.customers}
    invoices = {i.invoice_id: i for i in ds.invoices}
    for inv in ds.invoices:
        if inv.customer_id not in terms:
            problems.append(f"{inv.invoice_id}: unknown customer")
            continue
        if not RP_START <= inv.issue_date <= VD:
            problems.append(f"{inv.invoice_id}: issued outside RP")
        if inv.issue_date.weekday() not in ISSUE_WEEKDAYS:
            problems.append(f"{inv.invoice_id}: issued on Friday/Saturday")
        if inv.due_date != inv.issue_date + timedelta(days=terms[inv.customer_id]):
            problems.append(f"{inv.invoice_id}: due date != issue + terms")
        if inv.net <= 0 or inv.vat != vat_for(inv.net) or inv.total != inv.net + inv.vat:
            problems.append(f"{inv.invoice_id}: amounts inconsistent")

    settled = defaultdict(int)
    credit_count = defaultdict(int)
    links = [(p.payment_id, p.invoice_id, p.payment_date, p.amount) for p in ds.payments]
    links += [(c.credit_note_id, c.invoice_id, c.credit_date, c.total) for c in ds.credit_notes]
    for record_id, invoice_id, d, amount in links:
        inv = invoices.get(invoice_id)
        if inv is None:
            problems.append(f"{record_id}: unknown invoice")
            continue
        if amount <= 0:
            problems.append(f"{record_id}: amount not positive")
        if d < inv.issue_date:
            problems.append(f"{record_id}: dated before its invoice")
        if d > EXPORT_DATE:
            problems.append(f"{record_id}: dated after export date")
        settled[invoice_id] += amount
    for cn in ds.credit_notes:
        credit_count[cn.invoice_id] += 1
        if cn.net <= 0 or cn.vat != vat_for(cn.net) or cn.total != cn.net + cn.vat:
            problems.append(f"{cn.credit_note_id}: amounts inconsistent")
    for invoice_id, amount in settled.items():
        if invoice_id in invoices and amount > invoices[invoice_id].total:
            problems.append(f"{invoice_id}: payments + credit notes exceed total")
    problems += [f"{i}: more than one credit note" for i, n in credit_count.items() if n > 1]

    months_with_invoices = {(i.issue_date.year, i.issue_date.month) for i in ds.invoices}
    problems += [f"{y}-{m:02d}: no invoices" for y, m in rp_months() if (y, m) not in months_with_invoices]
    return problems


def profile_report(ds: Dataset) -> dict:
    """Generator self-check of the clean sample against the §3 targets.

    This summarises the generator's own in-memory data; it is not the analysis engine.
    """
    invoices = {i.invoice_id: i for i in ds.invoices}
    gross = sum(i.net for i in ds.invoices)
    credit_in_rp = [c for c in ds.credit_notes if c.credit_date <= VD]
    credit_net = sum(c.net for c in credit_in_rp)

    net_by_customer = defaultdict(int)
    for inv in ds.invoices:
        net_by_customer[inv.customer_id] += inv.net
    for cn in credit_in_rp:
        net_by_customer[invoices[cn.invoice_id].customer_id] -= cn.net
    net_sales = gross - credit_net
    shares = sorted((v / net_sales for v in net_by_customer.values()), reverse=True)

    settled = defaultdict(int)
    for p in ds.payments:
        if p.payment_date <= VD:
            settled[p.invoice_id] += p.amount
    for cn in credit_in_rp:
        settled[cn.invoice_id] += cn.total
    open_by_invoice = {k: inv.total - settled[k] for k, inv in invoices.items()}
    open_total = sum(open_by_invoice.values())
    over_90 = sum(v for k, v in open_by_invoice.items() if (VD - invoices[k].due_date).days > 90)
    monthly_billing = sum(i.total for i in ds.invoices) / 12

    states = defaultdict(int)
    for k, inv in invoices.items():
        states["fully_settled" if open_by_invoice[k] == 0
               else "unpaid" if open_by_invoice[k] == inv.total else "partly_settled"] += 1

    months = {(i.issue_date.year, i.issue_date.month) for i in ds.invoices}
    actual = {
        "gross_sales_sar": gross / 100,
        "invoice_count": len(ds.invoices),
        "top1_net_sales_share": round(shares[0], 4),
        "top5_net_sales_share": round(sum(shares[:5]), 4),
        "credit_note_ratio": round(credit_net / gross, 4),
        "open_balance_months_of_sales": round(open_total / monthly_billing, 2),
        "over_90_share_of_open_balance": round(over_90 / open_total, 4),
        "empty_months": sum(1 for ym in rp_months() if ym not in months),
    }
    return {
        "seed": SEED,
        "note": "Generator self-check of the clean sample; not the analysis engine.",
        "record_counts": {"customers": len(ds.customers), "invoices": len(ds.invoices),
                          "payments": len(ds.payments), "credit_notes": len(ds.credit_notes)},
        "invoice_settlement_at_vd": dict(sorted(states.items())),
        "open_balance_at_vd_sar": open_total / 100,
        "after_vd_records": {
            "payments": sum(1 for p in ds.payments if p.payment_date > VD),
            "credit_notes": sum(1 for c in ds.credit_notes if c.credit_date > VD),
        },
        "monthly_invoice_counts": {
            f"{y}-{m:02d}": sum(1 for i in ds.invoices if (i.issue_date.year, i.issue_date.month) == (y, m))
            for y, m in rp_months()
        },
        "profile_vs_targets": {
            key: {"actual": actual[key], "target": list(TARGETS[key]),
                  "within_target": TARGETS[key][0] <= actual[key] <= TARGETS[key][1]}
            for key in TARGETS
        },
    }


def clean_rows(ds: Dataset) -> dict[str, list[dict]]:
    return {
        "customers": [{"customer_id": c.customer_id, "customer_name": c.name, "city": "Riyadh",
                       "segment": c.segment, "payment_terms_days": c.terms} for c in ds.customers],
        "invoices": [{"invoice_id": i.invoice_id, "customer_id": i.customer_id,
                      "issue_date": i.issue_date, "due_date": i.due_date, "currency": "SAR",
                      "net_amount": i.net, "vat_amount": i.vat, "total_amount": i.total}
                     for i in ds.invoices],
        "payments": [{"payment_id": p.payment_id, "invoice_id": p.invoice_id,
                      "payment_date": p.payment_date, "amount": p.amount} for p in ds.payments],
        "credit_notes": [{"credit_note_id": c.credit_note_id, "invoice_id": c.invoice_id,
                          "credit_date": c.credit_date, "net_amount": c.net, "vat_amount": c.vat,
                          "total_amount": c.total, "reason": c.reason} for c in ds.credit_notes],
    }


# ---------------------------------------------------------------- record builders
# Amount fields take int halalas; any other value (text, None) is written as-is so
# that defects can be expressed exactly.

def cust(customer_id, name, segment="grocery", terms=30) -> dict:
    return {"customer_id": customer_id, "customer_name": name, "city": "Riyadh",
            "segment": segment, "payment_terms_days": terms}


def inv(invoice_id, customer_id, issue, due, net, vat, total, currency="SAR") -> dict:
    return {"invoice_id": invoice_id, "customer_id": customer_id, "issue_date": issue,
            "due_date": due, "currency": currency, "net_amount": net, "vat_amount": vat,
            "total_amount": total}


def pay(payment_id, invoice_id, payment_date, amount) -> dict:
    return {"payment_id": payment_id, "invoice_id": invoice_id,
            "payment_date": payment_date, "amount": amount}


def cn(credit_note_id, invoice_id, credit_date, net, vat, total, reason="return") -> dict:
    return {"credit_note_id": credit_note_id, "invoice_id": invoice_id, "credit_date": credit_date,
            "net_amount": net, "vat_amount": vat, "total_amount": total, "reason": reason}


# ---------------------------------------------------------------- Dataset A (§5.1)

def defect_rows(customers: tuple[Customer, ...]) -> dict[str, list[dict]]:
    name_007 = next(c.name for c in customers if c.customer_id == "CUST-007")

    def d_inv(invoice_id, customer_id="CUST-D20", issue=date(2026, 6, 10), due=date(2026, 7, 10),
              net=100_000, vat=15_000, total=115_000, currency="SAR"):
        return inv(invoice_id, customer_id, issue, due, net, vat, total, currency)

    return {
        "customers": [
            cust("CUST-D20", "Defect Test Store"),
            cust(None, "Missing ID Store"),                        # A-01
            cust("CUST-D21", "Duplicate Store A"),                 # A-02
            cust("CUST-D21", "Duplicate Store B"),                 # A-02
            cust("CUST-D22", name_007),                            # A-16
        ],
        "invoices": [
            d_inv("INV-D01", customer_id="CUST-D21"),              # A-03
            d_inv("INV-D02", customer_id="CUST-999"),              # A-04
            d_inv(None),                                           # A-05
            d_inv("INV-D03"),                                      # A-06
            d_inv("INV-D03", net=200_000, vat=30_000, total=230_000),  # A-06
            d_inv("INV-D04", net=0, vat=0, total=0),               # A-07
            d_inv("INV-D05", net="abc"),                           # A-08
            d_inv("INV-D06", due=date(2026, 6, 1)),                # A-09
            d_inv("INV-D07", currency="USD"),                      # A-10
            d_inv("INV-D08", total=115_002),                       # A-11
            d_inv("INV-D09", total=115_001),                       # A-12
            d_inv("INV-D10", issue=date(2025, 9, 20), due=date(2025, 10, 20)),  # A-13
            d_inv("INV-D11"),                                      # A-14
            d_inv("INV-D12"),                                      # A-15
        ],
        "payments": [
            pay("PAY-D03", "INV-D03", date(2026, 6, 20), 115_000),   # A-06
            pay("PAY-D10", "INV-D10", date(2025, 10, 15), 115_000),  # A-13
            pay("PAY-D11A", "INV-D11", date(2026, 6, 20), 115_000),  # A-14
            pay("PAY-D11B", "INV-D11", date(2026, 6, 25), 10_000),   # A-14
            pay("PAY-D12", "INV-D12", date(2026, 6, 5), 115_000),    # A-15
        ],
        "credit_notes": [],
    }


# ---------------------------------------------------------------- Dataset B (§5.2)

def case_sheets(customers=(), invoices=None, payments=(), credit_notes=()) -> dict[str, list[dict]]:
    base_invoices = [
        inv("INV-T1", "CUST-T1", date(2026, 8, 1), date(2026, 8, 31), 100_000, 15_000, 115_000),
        inv("INV-T2", "CUST-T1", date(2026, 9, 10), date(2026, 10, 10), 200_000, 30_000, 230_000),
    ]
    return {
        "customers": [cust("CUST-T1", "Test Customer One"), *customers],
        "invoices": list(base_invoices if invoices is None else invoices),
        "payments": list(payments),
        "credit_notes": list(credit_notes),
    }


def case_workbooks() -> dict[str, dict[str, list[dict]]]:
    """TC-02 .. TC-18 workbooks. TC-01 (CSV) and TC-02's missing column are handled by the writer."""
    return {
        "TC-02": case_sheets(),
        "TC-03": case_sheets(invoices=[
            inv("INV-T1", "CUST-T1", date(2025, 9, 1), date(2025, 10, 1), 100_000, 15_000, 115_000),
            inv("INV-T2", "CUST-T1", date(2025, 9, 10), date(2025, 10, 10), 200_000, 30_000, 230_000),
        ]),
        "TC-04": case_sheets(payments=[pay("PAY-T1", "INV-T1", date(2026, 8, 20), 0)]),
        "TC-05": case_sheets(payments=[pay("PAY-T1", "INV-T1", "2026-13-01", 115_000)]),
        "TC-06": case_sheets(payments=[pay(None, "INV-T1", date(2026, 8, 20), 50_000)]),
        "TC-07": case_sheets(payments=[pay("PAY-T1", "INV-T1", date(2026, 8, 20), 50_000),
                                       pay("PAY-T1", "INV-T2", date(2026, 9, 15), 50_000)]),
        "TC-08": case_sheets(payments=[pay("PAY-T1", "INV-X99", date(2026, 8, 20), 50_000)]),
        "TC-09": case_sheets(credit_notes=[cn("CN-T1", "INV-T1", date(2026, 8, 10), 10_000, "x", 11_500)]),
        "TC-10": case_sheets(credit_notes=[cn("CN-T1", "INV-T1", date(2026, 8, 10), 10_000, 1_500, 11_505)]),
        "TC-11": case_sheets(credit_notes=[cn("CN-T1", "INV-T1", None, 10_000, 1_500, 11_500)]),
        "TC-12": case_sheets(credit_notes=[cn("CN-T1", "INV-X99", date(2026, 8, 10), 10_000, 1_500, 11_500)]),
        "TC-13": case_sheets(payments=[pay("PAY-T1", "INV-T1", date(2026, 10, 5), 115_000)]),
        "TC-14": case_sheets(),
        "TC-15": case_sheets(payments=[pay("PAY-T1", "INV-T1", date(2026, 8, 25), 115_000),
                                       pay("PAY-T2", "INV-T2", date(2026, 9, 20), 230_000)]),
        "TC-16": case_sheets(
            customers=[cust("CUST-T2", "Test Customer Two")],
            invoices=[*case_sheets()["invoices"],
                      inv("INV-T3", "CUST-T2", date(2026, 9, 1), date(2026, 10, 1), 50_000, 7_500, 57_500)],
        ),
        "TC-17": case_sheets(invoices=[
            *case_sheets()["invoices"],
            inv("INV-T4", "CUST-T1", date(2026, 8, 31), date(2026, 9, 30), 10_000, 1_500, 11_500),
            inv("INV-T5", "CUST-T1", date(2026, 8, 30), date(2026, 9, 29), 10_000, 1_500, 11_500),
        ]),
        "TC-18": case_sheets(credit_notes=[cn("CN-T1", "INV-T1", date(2026, 10, 5), 10_000, 1_500, 11_500)]),
    }


# ---------------------------------------------------------------- worked example (§6)

def worked_example_sheets() -> dict[str, list[dict]]:
    return {
        "customers": [cust("CUST-A", "Worked Customer A", terms=30),
                      cust("CUST-B", "Worked Customer B", terms=60)],
        "invoices": [
            inv("W1", "CUST-A", date(2026, 6, 15), date(2026, 7, 15), 1_000_000, 150_000, 1_150_000),
            inv("W2", "CUST-A", date(2026, 8, 1), date(2026, 8, 31), 400_000, 60_000, 460_000),
            inv("W3", "CUST-B", date(2026, 8, 1), date(2026, 9, 30), 600_000, 90_000, 690_000),
            inv("W4", "CUST-B", date(2026, 9, 10), date(2026, 11, 9), 200_033, 30_005, 230_038),
            inv("W5", "CUST-A", date(2026, 5, 10), date(2026, 6, 9), 300_000, 45_000, 345_000),
        ],
        "payments": [
            pay("P1", "W1", date(2026, 7, 10), 500_000),
            pay("P2", "W1", date(2026, 8, 20), 300_000),
            pay("P3", "W2", date(2026, 9, 5), 400_000),
            pay("P4", "W3", date(2026, 10, 5), 690_000),
        ],
        "credit_notes": [cn("CN1", "W2", date(2026, 8, 10), 40_000, 6_000, 46_000)],
    }


# ---------------------------------------------------------------- writers

def cell_value(column: str, value):
    if column in AMOUNT_COLUMNS and isinstance(value, int):
        return halalas_to_cell(value)
    return value


def write_workbook(path: Path, sheets: dict[str, list[dict]], drop_columns=frozenset()) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, columns in COLUMNS.items():
        ws = wb.create_sheet(sheet_name)
        columns = [c for c in columns if f"{sheet_name}.{c}" not in drop_columns]
        ws.append(columns)
        for row in sheets.get(sheet_name, []):
            ws.append([cell_value(c, row.get(c)) for c in columns])
            for cell, column in zip(ws[ws.max_row], columns):
                if column in AMOUNT_COLUMNS and isinstance(cell.value, float):
                    cell.number_format = "0.00"
                elif column in DATE_COLUMNS and isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    def text(column, value):
        if column in AMOUNT_COLUMNS and isinstance(value, int):
            return f"{value // 100}.{value % 100:02d}"
        return value.isoformat() if isinstance(value, date) else value

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        writer.writerows([[text(c, r[c]) for c in columns] for r in rows])


def write_outputs(sample_dir: Path = SAMPLE_DIR, fixtures_dir: Path = FIXTURES_DIR,
                  seed: int = SEED) -> dict:
    """Generate everything and return the clean-sample profile report."""
    ds = generate_clean(seed)
    rows = clean_rows(ds)
    write_workbook(sample_dir / "sample_clean.xlsx", rows)

    defects = defect_rows(ds.customers)
    write_workbook(sample_dir / "sample_defects.xlsx",
                   {sheet: rows[sheet] + defects[sheet] for sheet in COLUMNS})

    report = profile_report(ds)
    report["seed"] = seed
    (sample_dir / "generation_report.json").write_text(
        json.dumps(report, indent=2) + "\n", encoding="utf-8")

    cases_dir = fixtures_dir / "cases"
    write_csv(cases_dir / "TC-01.csv", case_sheets()["invoices"], COLUMNS["invoices"])
    for name, sheets in case_workbooks().items():
        drop = {"invoices.due_date"} if name == "TC-02" else frozenset()
        write_workbook(cases_dir / f"{name}.xlsx", sheets, drop)
    write_workbook(fixtures_dir / "worked_example.xlsx", worked_example_sheets())
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args()
    report = write_outputs(seed=args.seed)
    print(json.dumps({"record_counts": report["record_counts"],
                      "profile_vs_targets": report["profile_vs_targets"]}, indent=2))


if __name__ == "__main__":
    main()
