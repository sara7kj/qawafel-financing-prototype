"""Generator integrity, Excel amount round trip and data-content stability.

These tests check the generator and the fixtures it writes. They do not run any
analysis logic; the expected-results JSON files are hand-written and only
checked here for arithmetic consistency.
"""

import csv
import json
import random
from collections import defaultdict
from dataclasses import replace
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts import generate_invoices as gen

EXPECTED = gen.FIXTURES_DIR / "expected"
GENERATED_FILES = [
    Path("sample") / "sample_clean.xlsx",
    Path("sample") / "sample_defects.xlsx",
    Path("fixtures") / "worked_example.xlsx",
    *[Path("fixtures") / "cases" / f"TC-{n:02d}.xlsx" for n in range(2, 19)],
]


def cell_to_halalas(value) -> int:
    """Test-only read-back: the shortest decimal of the stored double, in halalas."""
    amount = Decimal(str(value))
    assert amount == amount.quantize(Decimal("0.01")), f"{value!r} has more than 2 decimals"
    return int(amount * 100)


def workbook_content(path: Path) -> dict[str, list[tuple]]:
    wb = load_workbook(path)
    return {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def committed(rel: Path) -> Path:
    return (gen.SAMPLE_DIR.parent if rel.parts[0] == "sample" else gen.FIXTURES_DIR.parent) / rel


def sar(text: str) -> Decimal:
    return Decimal(text)


@pytest.fixture(scope="module")
def clean():
    return gen.generate_clean()


# ---------------------------------------------------------------- amounts

def test_amounts_round_trip_through_xlsx(tmp_path):
    rng = random.Random(1)
    values = [1, 5, 10, 33, 99, 100, 101, 30_005, 115_001, 115_002, 200_033, 999_999_999]
    values += list(range(1, 10_001))                                   # every halala up to SAR 100
    values += [rng.randint(1, 1_000_000_000) for _ in range(20_000)]  # up to SAR 10 million
    wb = Workbook()
    for v in values:
        wb.active.append([gen.halalas_to_cell(v)])
    path = tmp_path / "amounts.xlsx"
    wb.save(path)

    read_back = [row[0] for row in load_workbook(path).active.iter_rows(values_only=True)]
    assert [cell_to_halalas(v) for v in read_back] == values


def test_sample_amounts_recover_exact_halalas(clean):
    content = workbook_content(gen.SAMPLE_DIR / "sample_clean.xlsx")
    invoices = content["invoices"][1:]
    assert [(cell_to_halalas(r[5]), cell_to_halalas(r[6]), cell_to_halalas(r[7])) for r in invoices] == \
        [(i.net, i.vat, i.total) for i in clean.invoices]
    assert [cell_to_halalas(r[3]) for r in content["payments"][1:]] == [p.amount for p in clean.payments]
    assert [cell_to_halalas(r[5]) for r in content["credit_notes"][1:]] == [c.total for c in clean.credit_notes]


def test_vat_rounds_half_up():
    assert gen.vat_for(200_033) == 30_005   # 300.0495 -> 300.05
    assert gen.vat_for(10) == 2             # 0.015 -> 0.02
    assert gen.vat_for(100_000) == 15_000


# ---------------------------------------------------------------- clean sample integrity

def test_clean_sample_passes_generator_checks(clean):
    assert gen.check_clean(clean) == []


def test_clean_settlements_never_exceed_invoice_total(clean):
    settled = defaultdict(int)
    for p in clean.payments:
        settled[p.invoice_id] += p.amount
    for c in clean.credit_notes:
        settled[c.invoice_id] += c.total
    totals = {i.invoice_id: i.total for i in clean.invoices}
    assert all(amount <= totals[invoice_id] for invoice_id, amount in settled.items())


def test_clean_sample_shape(clean):
    assert len(clean.customers) == 25
    assert all(gen.RP_START <= i.issue_date <= gen.VD for i in clean.invoices)
    months = {(i.issue_date.year, i.issue_date.month) for i in clean.invoices}
    assert months == set(gen.rp_months())
    issue = {i.invoice_id: i.issue_date for i in clean.invoices}
    links = [(p.invoice_id, p.payment_date) for p in clean.payments]
    links += [(c.invoice_id, c.credit_date) for c in clean.credit_notes]
    assert all(issue[i] <= d <= gen.EXPORT_DATE for i, d in links)

    payment_counts = defaultdict(int)
    for p in clean.payments:
        payment_counts[p.invoice_id] += 1
    assert any(n > 1 for n in payment_counts.values()), "no partial payments"
    assert any(i.invoice_id not in payment_counts for i in clean.invoices), "no unpaid invoices"
    assert clean.credit_notes, "no credit notes"
    assert any(p.payment_date > gen.VD for p in clean.payments), "no settlements after VD"


def test_check_clean_detects_over_settlement(clean):
    inv = clean.invoices[0]
    extra = gen.Payment("PAY-X", inv.invoice_id, inv.issue_date, inv.total + 1)
    broken = replace(clean, payments=clean.payments + (extra,))
    assert any("exceed" in p for p in gen.check_clean(broken))


# ---------------------------------------------------------------- reproducibility

def test_same_seed_gives_identical_content(tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    gen.write_outputs(a / "sample", a / "fixtures")
    gen.write_outputs(b / "sample", b / "fixtures")
    for rel in GENERATED_FILES:
        assert workbook_content(a / rel) == workbook_content(b / rel), rel
    for rel in [Path("sample") / "generation_report.json", Path("fixtures") / "cases" / "TC-01.csv"]:
        assert (a / rel).read_text(encoding="utf-8") == (b / rel).read_text(encoding="utf-8"), rel


def test_different_seed_changes_content(clean):
    assert gen.generate_clean(gen.SEED + 1).invoices != clean.invoices


def test_committed_outputs_match_generator(tmp_path):
    gen.write_outputs(tmp_path / "sample", tmp_path / "fixtures")
    for rel in GENERATED_FILES:
        assert workbook_content(tmp_path / rel) == workbook_content(committed(rel)), \
            f"{rel} is out of date; rerun scripts/generate_invoices.py"
    for rel in [Path("sample") / "generation_report.json", Path("fixtures") / "cases" / "TC-01.csv"]:
        assert (tmp_path / rel).read_text(encoding="utf-8") == committed(rel).read_text(encoding="utf-8"), rel


# ---------------------------------------------------------------- Dataset A

def test_defect_workbook_is_clean_plus_dataset_a():
    clean_content = workbook_content(gen.SAMPLE_DIR / "sample_clean.xlsx")
    defect_content = workbook_content(gen.SAMPLE_DIR / "sample_defects.xlsx")
    added = {}
    for sheet, rows in clean_content.items():
        assert defect_content[sheet][:len(rows)] == rows, f"clean rows changed in {sheet}"
        added[sheet] = defect_content[sheet][len(rows):]

    assert [r[0] for r in added["customers"]] == ["CUST-D20", None, "CUST-D21", "CUST-D21", "CUST-D22"]
    assert [r[0] for r in added["invoices"]] == [
        "INV-D01", "INV-D02", None, "INV-D03", "INV-D03", "INV-D04", "INV-D05", "INV-D06",
        "INV-D07", "INV-D08", "INV-D09", "INV-D10", "INV-D11", "INV-D12"]
    assert [r[0] for r in added["payments"]] == ["PAY-D03", "PAY-D10", "PAY-D11A", "PAY-D11B", "PAY-D12"]
    assert added["credit_notes"] == []

    invoices = {r[0]: r for r in added["invoices"] if r[0] not in (None, "INV-D03")}
    assert invoices["INV-D05"][5] == "abc" and cell_to_halalas(invoices["INV-D05"][7]) == 115_000
    assert invoices["INV-D07"][4] == "USD"
    assert cell_to_halalas(invoices["INV-D08"][7]) == 115_002
    assert cell_to_halalas(invoices["INV-D09"][7]) == 115_001
    assert invoices["INV-D06"][3] == datetime(2026, 6, 1)
    names = {r[0]: r[1] for r in clean_content["customers"][1:] + added["customers"]}
    assert names["CUST-D22"] == names["CUST-007"] == "Reem Mini Market 07"


# ---------------------------------------------------------------- Dataset B and worked example

def test_every_case_has_fixture_and_expected_entry():
    expected = json.loads((EXPECTED / "cases_expected.json").read_text(encoding="utf-8"))
    case_ids = {k for k in expected if k.startswith("TC-")}
    assert case_ids == {f"TC-{n:02d}" for n in range(1, 19)}
    for case_id in case_ids:
        assert (gen.FIXTURES_DIR / expected[case_id]["file"]).exists(), case_id


def test_case_fixtures_contain_the_designed_defects():
    cases = gen.FIXTURES_DIR / "cases"
    with (cases / "TC-01.csv").open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[1] == ["INV-T1", "CUST-T1", "2026-08-01", "2026-08-31", "SAR", "1000.00", "150.00", "1150.00"]

    def sheet(case, name):
        return workbook_content(cases / f"{case}.xlsx")[name]

    assert "due_date" not in sheet("TC-02", "invoices")[0]
    assert [r[2] for r in sheet("TC-03", "invoices")[1:]] == [datetime(2025, 9, 1), datetime(2025, 9, 10)]
    assert sheet("TC-04", "payments")[1][3] == 0
    assert sheet("TC-05", "payments")[1][2] == "2026-13-01"
    assert sheet("TC-06", "payments")[1][0] is None
    assert [r[0] for r in sheet("TC-07", "payments")[1:]] == ["PAY-T1", "PAY-T1"]
    assert sheet("TC-08", "payments")[1][1] == "INV-X99"
    assert sheet("TC-09", "credit_notes")[1][4] == "x"
    assert cell_to_halalas(sheet("TC-10", "credit_notes")[1][5]) == 11_505
    assert sheet("TC-11", "credit_notes")[1][2] is None
    assert sheet("TC-12", "credit_notes")[1][1] == "INV-X99"
    assert sheet("TC-13", "payments")[1][2] == datetime(2026, 10, 5)
    assert len(sheet("TC-14", "payments")) == 1 and len(sheet("TC-14", "credit_notes")) == 1
    assert [r[0] for r in sheet("TC-16", "invoices")[1:]] == ["INV-T1", "INV-T2", "INV-T3"]
    assert [r[3] for r in sheet("TC-17", "invoices")[3:]] == [datetime(2026, 9, 30), datetime(2026, 9, 29)]


def _monthly_sum(series):
    return sum((sar(v) for v in series["monthly"].values()), Decimal("0.00"))


def _financing(eligible: str) -> Decimal:
    return (sar(eligible) * Decimal("0.70")).quantize(Decimal("0.01"), ROUND_HALF_UP)


def test_expected_case_values_are_arithmetically_consistent():
    expected = json.loads((EXPECTED / "cases_expected.json").read_text(encoding="utf-8"))
    for case_id, case in expected.items():
        if not case_id.startswith("TC-") or case["block"]:
            continue
        assert _monthly_sum(case["gross_sales"]) == sar(case["gross_sales"]["total"]), case_id
        if case["collections"]["total"] is not None:
            assert _monthly_sum(case["collections"]) == sar(case["collections"]["total"]), case_id
        if case["aging"] is not None:
            buckets = [sar(v) for k, v in case["aging"].items() if k != "total"]
            assert sum(buckets) == sar(case["aging"]["total"]), case_id
        if case["financing_estimate_at_70pct"] is not None:
            assert _financing(case["eligible_receivables"]) == sar(case["financing_estimate_at_70pct"]), case_id


def test_worked_example_workbook_matches_design():
    content = workbook_content(gen.FIXTURES_DIR / "worked_example.xlsx")
    invoices = [(r[0], r[1], r[2].date(), r[3].date(), cell_to_halalas(r[5]), cell_to_halalas(r[6]),
                 cell_to_halalas(r[7])) for r in content["invoices"][1:]]
    assert invoices == [
        ("W1", "CUST-A", datetime(2026, 6, 15).date(), datetime(2026, 7, 15).date(), 1_000_000, 150_000, 1_150_000),
        ("W2", "CUST-A", datetime(2026, 8, 1).date(), datetime(2026, 8, 31).date(), 400_000, 60_000, 460_000),
        ("W3", "CUST-B", datetime(2026, 8, 1).date(), datetime(2026, 9, 30).date(), 600_000, 90_000, 690_000),
        ("W4", "CUST-B", datetime(2026, 9, 10).date(), datetime(2026, 11, 9).date(), 200_033, 30_005, 230_038),
        ("W5", "CUST-A", datetime(2026, 5, 10).date(), datetime(2026, 6, 9).date(), 300_000, 45_000, 345_000),
    ]
    payments = [(r[0], r[1], r[2].date(), cell_to_halalas(r[3])) for r in content["payments"][1:]]
    assert payments == [
        ("P1", "W1", datetime(2026, 7, 10).date(), 500_000),
        ("P2", "W1", datetime(2026, 8, 20).date(), 300_000),
        ("P3", "W2", datetime(2026, 9, 5).date(), 400_000),
        ("P4", "W3", datetime(2026, 10, 5).date(), 690_000),
    ]
    credit = content["credit_notes"][1]
    assert (credit[0], credit[1], credit[2].date(), cell_to_halalas(credit[5])) == \
        ("CN1", "W2", datetime(2026, 8, 10).date(), 46_000)


def test_worked_example_expected_values_are_arithmetically_consistent():
    exp = json.loads((EXPECTED / "worked_example_expected.json").read_text(encoding="utf-8"))
    for series in ("gross_sales", "credit_notes", "net_sales", "collections_from_accepted_records_only"):
        assert _monthly_sum(exp[series]) == sar(exp[series]["total"]), series
    assert sum(sar(v["open"]) for v in exp["open_balances"].values()) == sar(exp["aging"]["total"])
    assert sum(sar(v) for k, v in exp["aging"].items() if k != "total") == sar(exp["aging"]["total"])
    assert _financing(exp["eligible_receivables"]) == sar(exp["financing_estimate_at_70pct"])
