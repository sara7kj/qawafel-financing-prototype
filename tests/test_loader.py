"""Loader: sources, structure checks (S-01, S-02), row numbers, raw values and cell parsing."""

import io
import math
from datetime import date, datetime
from decimal import Decimal

import pytest

from conftest import FIXTURES, write_raw_workbook
from financing.loader import load_workbook_file, parse_amount, parse_date, parse_id
from financing.schema import Severity

CUSTOMER = {"customer_id": "C1", "customer_name": "One", "payment_terms_days": 30}
INVOICE = {"invoice_id": "I1", "customer_id": "C1", "issue_date": date(2026, 8, 1),
           "due_date": date(2026, 8, 31), "currency": "SAR",
           "net_amount": 1000.0, "vat_amount": 150.0, "total_amount": 1150.0}


def block_rules(result):
    return [i.rule for i in result.issues if i.severity is Severity.BLOCK]


# ---------------------------------------------------------------- amounts

@pytest.mark.parametrize("value, halalas", [
    (5, 500),
    (1150.01, 115001),
    (2000.33, 200033),
    (0.01, 1),
    (17.2500000001, 1725),          # float noise is absorbed
    (" 12.50 ", 1250),
    ("1150", 115000),
    ("-5", -500),                   # parsed; the >0 rule is validation's job
    (0, 0),
    (0.0, 0),
    (Decimal("10.10"), 1010),
])
def test_parse_amount_accepts(value, halalas):
    assert parse_amount(value) == (halalas, None)


@pytest.mark.parametrize("value, kind", [
    (None, "missing"),
    ("   ", "missing"),
    (10.005, "invalid"),            # a third decimal is never rounded away
    ("10.005", "invalid"),
    (Decimal("1.234"), "invalid"),
    ("1,000.00", "invalid"),
    ("SAR 10", "invalid"),
    ("1e3", "invalid"),
    ("abc", "invalid"),
    (math.nan, "invalid"),
    (math.inf, "invalid"),
    (-math.inf, "invalid"),
    ("NaN", "invalid"),
    ("Infinity", "invalid"),
    (Decimal("NaN"), "invalid"),
    (True, "invalid"),
    (False, "invalid"),
    (datetime(2026, 1, 1), "invalid"),
])
def test_parse_amount_rejects(value, kind):
    parsed, error = parse_amount(value)
    assert parsed is None and error[0] == kind


# ---------------------------------------------------------------- dates and IDs

@pytest.mark.parametrize("value, expected", [
    (datetime(2026, 8, 1, 13, 45), date(2026, 8, 1)),
    (date(2026, 8, 1), date(2026, 8, 1)),
    ("2026-08-01", date(2026, 8, 1)),
    (" 2026-08-01 ", date(2026, 8, 1)),
])
def test_parse_date_accepts(value, expected):
    assert parse_date(value) == (expected, None)


@pytest.mark.parametrize("value", ["2026-13-01", "01/08/2026", "2026-8-1", 46000, 46000.0, True])
def test_parse_date_rejects(value):
    assert parse_date(value)[0] is None and parse_date(value)[1][0] == "invalid"


@pytest.mark.parametrize("value, expected", [
    (" INV-1 ", "INV-1"), (1001, "1001"), (1001.0, "1001"), ("", None), ("  ", None), (None, None),
])
def test_parse_id(value, expected):
    assert parse_id(value)[0] == expected


# ---------------------------------------------------------------- sources

def test_load_from_path_bytes_and_uploaded_file_agree(tmp_path):
    path = write_raw_workbook(tmp_path / "book.xlsx", {"customers": [CUSTOMER], "invoices": [INVOICE]})
    data = path.read_bytes()
    uploaded = io.BytesIO(data)
    uploaded.name = "upload.xlsx"  # like a Streamlit UploadedFile
    results = [load_workbook_file(path), load_workbook_file(data), load_workbook_file(uploaded),
               load_workbook_file(io.BytesIO(data))]
    for result in results:
        assert not result.blocked
        assert [(r.row, r.values) for r in result.rows["invoices"]] == \
            [(r.row, r.values) for r in results[0].rows["invoices"]]
    assert results[2].source_name == "upload.xlsx"


def test_row_numbers_and_raw_values_are_kept(tmp_path):
    bad = {**INVOICE, "invoice_id": "I2", "net_amount": "abc"}
    path = write_raw_workbook(tmp_path / "book.xlsx",
                              {"customers": [CUSTOMER], "invoices": [INVOICE, {}, bad]})
    records = load_workbook_file(path).rows["invoices"]
    assert [r.row for r in records] == [2, 4]  # the empty Excel row 3 is skipped, numbering is kept
    assert records[1].raw["net_amount"] == "abc"
    assert records[1].values["net_amount"] is None
    assert records[1].errors["net_amount"].kind == "invalid"
    assert records[1].values["total_amount"] == 115000


# ---------------------------------------------------------------- structure

def test_csv_is_blocked_s01():
    assert block_rules(load_workbook_file(FIXTURES / "cases" / "TC-01.csv")) == ["S-01"]


def test_wrong_extension_and_unreadable_bytes_are_blocked_s01(tmp_path):
    path = write_raw_workbook(tmp_path / "book.xlsx", {"customers": [CUSTOMER]})
    assert block_rules(load_workbook_file(path.read_bytes(), filename="book.xls")) == ["S-01"]
    garbage = io.BytesIO(b"not a workbook")
    garbage.name = "book.xlsx"
    assert block_rules(load_workbook_file(garbage)) == ["S-01"]


def test_missing_column_is_blocked_s02_and_named():
    result = load_workbook_file(FIXTURES / "cases" / "TC-02.xlsx")
    assert [(i.rule, i.sheet, i.column) for i in result.issues] == [("S-02", "invoices", "due_date")]


def test_missing_sheet_is_blocked_s02(tmp_path):
    path = write_raw_workbook(tmp_path / "book.xlsx", {"customers": [CUSTOMER]}, omit_sheets=("credit_notes",))
    result = load_workbook_file(path)
    assert [(i.rule, i.sheet) for i in result.issues] == [("S-02", "credit_notes")]


def _write_header_and_row(path, sheet, header, row):
    from openpyxl import Workbook
    from financing.schema import SHEETS
    wb = Workbook()
    wb.remove(wb.active)
    for name, spec in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(header if name == sheet else list(spec.columns))
        if name == sheet:
            ws.append(row)
    wb.save(path)
    return path


def test_conflicting_duplicate_amount_columns_are_blocked_s02(tmp_path):
    path = _write_header_and_row(tmp_path / "book.xlsx", "payments",
                                 ["payment_id", "invoice_id", "payment_date", "amount", "amount"],
                                 ["P1", "I1", date(2026, 8, 20), 100.0, 900.0])
    result = load_workbook_file(path)
    assert result.blocked
    [issue] = result.issues
    assert (issue.rule, issue.sheet, issue.column) == ("S-02", "payments", "amount")
    assert "columns D, E" in issue.message


def test_duplicate_id_column_is_blocked_s02_after_trimming(tmp_path):
    path = _write_header_and_row(tmp_path / "book.xlsx", "invoices",
                                 ["invoice_id", "customer_id", " invoice_id ", "issue_date", "due_date",
                                  "currency", "net_amount", "vat_amount", "total_amount"],
                                 ["I1", "C1", "I9", date(2026, 8, 1), date(2026, 8, 31), "SAR", 1000.0, 150.0, 1150.0])
    result = load_workbook_file(path)
    assert [(i.rule, i.sheet, i.column) for i in result.issues] == [("S-02", "invoices", "invoice_id")]
    assert "columns A, C" in result.issues[0].message


def test_repeated_unknown_columns_are_still_ignored(tmp_path):
    path = _write_header_and_row(tmp_path / "book.xlsx", "payments",
                                 ["payment_id", "note", "invoice_id", "payment_date", "amount", "note"],
                                 ["P1", "a", "I1", date(2026, 8, 20), 100.0, "b"])
    assert not load_workbook_file(path).blocked


def test_optional_columns_may_be_missing_and_extra_columns_are_ignored(tmp_path):
    columns = {"customers": ["notes", "customer_id", "customer_name", "payment_terms_days"],
               "credit_notes": ["credit_note_id", "invoice_id", "credit_date",
                                "net_amount", "vat_amount", "total_amount"]}
    path = write_raw_workbook(tmp_path / "book.xlsx",
                              {"customers": [{**CUSTOMER, "notes": "ignored"}]}, columns=columns)
    result = load_workbook_file(path)
    assert not result.blocked
    assert "notes" not in result.rows["customers"][0].raw
