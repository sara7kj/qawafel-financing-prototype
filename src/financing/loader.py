"""Load an invoice workbook into typed records.

Keeps every record's Excel row number and original cell values so that problems
can be explained. Only structural problems (S-01, S-02) are decided here; record,
reference and period rules live in validation.py.

Parsing rules: docs/data-design.md §1-§2. Amounts become integer halalas.
"""

from __future__ import annotations

import io
import math
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO, Callable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from financing.schema import SHEETS, FieldError, Issue, LoadResult, Record, Severity, SheetSpec

Parsed = tuple[Any, tuple[str, str] | None]  # (value, (error kind, detail) or None)

FLOAT_NOISE = Decimal("0.000001")
_DECIMAL_TEXT = re.compile(r"[+-]?\d+(?:\.(\d+))?")
_DATE_TEXT = re.compile(r"\d{4}-\d{2}-\d{2}")
_ID_COLUMNS = {"customer_id", "invoice_id", "payment_id", "credit_note_id"}


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_amount(value: Any) -> Parsed:
    """Cell value -> integer halalas. Never rounds beyond float noise (data-design §2)."""
    if _blank(value):
        return None, ("missing", "empty cell")
    if isinstance(value, bool):  # bool is a subclass of int
        return None, ("invalid", "a boolean is not an amount")
    if isinstance(value, int):
        return value * 100, None
    if isinstance(value, float):
        if not math.isfinite(value):
            return None, ("invalid", "NaN or infinity is not an amount")
        amount = Decimal(repr(value))
    elif isinstance(value, Decimal):
        if not value.is_finite():
            return None, ("invalid", "NaN or infinity is not an amount")
        amount = value
    elif isinstance(value, str):
        text = value.strip()
        match = _DECIMAL_TEXT.fullmatch(text)
        if not match:
            return None, ("invalid", "not a plain decimal number")
        if match.group(1) and len(match.group(1)) > 2:
            return None, ("invalid", "more than 2 decimal places")
        amount = Decimal(text)
    else:
        return None, ("invalid", f"a {type(value).__name__} cell is not an amount")

    halalas = amount.scaleb(2).to_integral_value()
    if abs(amount.scaleb(2) - halalas) > FLOAT_NOISE.scaleb(2):
        return None, ("invalid", "more than 2 decimal places")
    return int(halalas), None


def parse_date(value: Any) -> Parsed:
    """Excel date cell or ISO 'YYYY-MM-DD' text; any time part is ignored."""
    if _blank(value):
        return None, ("missing", "empty cell")
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, str) and _DATE_TEXT.fullmatch(value.strip()):
        try:
            return date.fromisoformat(value.strip()), None
        except ValueError:
            pass
    return None, ("invalid", "not a date cell or YYYY-MM-DD text")


def parse_id(value: Any) -> Parsed:
    """Trimmed text; integral numbers become text without decimals; blank is missing."""
    if isinstance(value, bool):
        return None, ("invalid", "a boolean is not an ID")
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = "" if value is None else str(value).strip()
    if not text:
        return None, ("missing", "empty cell")
    return text, None


def parse_int(value: Any) -> Parsed:
    if _blank(value):
        return None, ("missing", "empty cell")
    if isinstance(value, int) and not isinstance(value, bool):
        return value, None
    if isinstance(value, float) and value.is_integer():
        return int(value), None
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip()), None
    return None, ("invalid", "not a whole number")


def parse_text(value: Any) -> Parsed:
    if _blank(value):
        return None, None
    return str(value).strip(), None


def _parser(spec: SheetSpec, column: str) -> Callable[[Any], Parsed]:
    if column in spec.amount_columns:
        return parse_amount
    if column in spec.date_columns:
        return parse_date
    if column in _ID_COLUMNS:
        return parse_id
    if column == "payment_terms_days":
        return parse_int
    return parse_text


def parse_record(spec: SheetSpec, row: int, raw: dict[str, Any]) -> Record:
    values, errors = {}, {}
    for column, value in raw.items():
        parsed, error = _parser(spec, column)(value)
        values[column] = parsed
        if error:
            errors[column] = FieldError(column, error[0], error[1], value)
    return Record(spec.name, row, raw, values, errors)


def _source_name(source: Any) -> str | None:
    if isinstance(source, (str, Path)):
        return Path(source).name
    return getattr(source, "name", None)


def _open(source: Any) -> Any:
    if isinstance(source, (str, Path)):
        return str(source)
    if isinstance(source, (bytes, bytearray)):
        return io.BytesIO(source)
    return io.BytesIO(source.read())  # file-like, e.g. an uploaded file


def _block(rule: str, message: str, **kwargs) -> Issue:
    return Issue(rule, Severity.BLOCK, message, **kwargs)


def load_workbook_file(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> LoadResult:
    """Load from a path, raw bytes or a file-like object (e.g. an uploaded file).

    `filename` overrides the name used for the .xlsx extension check.
    """
    name = filename or _source_name(source)
    if name is not None and Path(name).suffix.lower() != ".xlsx":
        return LoadResult({}, [_block("S-01", f"'{name}' is not an .xlsx file")], name)
    try:
        wb = load_workbook(_open(source), read_only=True, data_only=True)
    except Exception as exc:  # any failure to open means the file is unreadable (S-01)
        return LoadResult({}, [_block("S-01", f"the file could not be read as an Excel workbook ({exc})")], name)

    rows: dict[str, list[Record]] = {}
    issues: list[Issue] = []
    try:
        for spec in SHEETS.values():
            if spec.name not in wb.sheetnames:
                issues.append(_block("S-02", f"required sheet '{spec.name}' is missing", sheet=spec.name))
                continue
            sheet_rows = wb[spec.name].iter_rows(values_only=True)
            header = next(sheet_rows, None) or ()
            found: dict[str, list[int]] = {}
            for index, cell in enumerate(header):
                column = cell.strip() if isinstance(cell, str) else cell
                if column in spec.columns:  # unknown extra columns are ignored
                    found.setdefault(column, []).append(index)
            missing = [c for c in spec.required if c not in found]
            for column in missing:
                issues.append(_block("S-02", f"required column '{spec.name}.{column}' is missing",
                                     sheet=spec.name, column=column))
            duplicated = {c: idx for c, idx in found.items() if len(idx) > 1}
            for column, indexes in duplicated.items():
                letters = ", ".join(get_column_letter(i + 1) for i in indexes)
                issues.append(_block("S-02", f"column '{spec.name}.{column}' appears more than once "
                                             f"(columns {letters}); the intended values are ambiguous",
                                     sheet=spec.name, column=column))
            if missing or duplicated:
                continue
            positions = {c: idx[0] for c, idx in found.items()}
            records = []
            for excel_row, cells in enumerate(sheet_rows, start=2):
                raw = {c: cells[i] if i < len(cells) else None for c, i in positions.items()}
                if all(_blank(v) for v in raw.values()):
                    continue  # fully empty rows are not records
                records.append(parse_record(spec, excel_row, raw))
            rows[spec.name] = records
    finally:
        wb.close()
    return LoadResult(rows, issues, name)
